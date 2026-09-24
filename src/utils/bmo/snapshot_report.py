"""A shift-report style Word document, generated from a snapshot.

NOT A SCREENSHOT. The page hides most of its content behind tabs and expanders,
and a picture of it would show whichever tab happened to be open. This builds the
report from the snapshot's DATA, so every section is written out in full and in
the same order every time:

    cover            headline KPIs and an executive summary in plain sentences
    blend            donut chart and a per-ore table (tonnes, Fe, slag, cost)
    constraints      every limit the optimiser was held to, met or violated
    fuel and coke    model anchor -> physics correction -> reported coke, as a waterfall
    slag             chemistry against its window, and where the slag comes from
    cost             ore / fuel / flux for each option
    options          LP, DE and the current blend side by side
    path, commentary when the build records them (UAT)
    plant data       what the run read, and whether it was frozen with the snapshot
    appendices       every input, optimiser messages, provenance

Sections with nothing to report are left out and the rest renumbered. Charts are
drawn with matplotlib (Plotly export needs ``kaleido``, which is not installed).
"""

from __future__ import annotations

import io
import math
from datetime import datetime
from typing import Any, Mapping

from utils.bmo.snapshot import _frame_rows, decode, ore_names, recommended_result

TEAL = "0F766E"
TEAL_LIGHT = "E6F4F1"
GREY_TEXT = "6B7280"
GREEN_FILL, GREEN_TEXT = "D1FAE5", "065F46"
RED_FILL, RED_TEXT = "FEE2E2", "991B1B"
AMBER_FILL = "FEF3C7"
ZEBRA = "F7F9FA"
CHART_COLOURS = ["#0F766E", "#14B8A6", "#F59E0B", "#6366F1", "#EF4444", "#84CC16",
                 "#EC4899", "#0EA5E9", "#A855F7", "#64748B", "#F97316", "#22C55E"]

# Input suffix -> (label, format). Anything not listed still appears, with a
# label derived from its key, so a new page input is never dropped from a report.
INPUT_LABELS: dict[str, tuple[str, str]] = {
    "target_production_mt": ("Target hot metal (MT)", "{:,.1f}"),
    "target_slag_rate_kg_per_thm": ("Max slag rate, plant basis (kg/THM)", "{:,.1f}"),
    "target_slag_basicity_min": ("B2 min (CaO/SiO2)", "{:.3f}"),
    "target_slag_basicity_max": ("B2 max (CaO/SiO2)", "{:.3f}"),
    "target_slag_t_basicity_min": ("T-basicity min", "{:.3f}"),
    "target_slag_t_basicity_max": ("T-basicity max", "{:.3f}"),
    "target_slag_al2o3_max_pct": ("Slag Al2O3 max (%)", "{:.2f}"),
    "target_slag_mgo_min_pct": ("Slag MgO min (%)", "{:.2f}"),
    "target_slag_mgo_al2o3_ratio_min": ("MgO/Al2O3 min", "{:.3f}"),
    "max_charges_per_hour": ("Max charges per hour", "{:.2f}"),
    "charge_mass_mt": ("Max qty per charge (MT)", "{:.2f}"),
    "burden_capacity_enabled": ("Enforce charging capacity", "{}"),
    "chemistry_mode": ("Chemistry mode", "{}"),
    "chemistry_window_days": ("Chemistry window (days)", "{}"),
    "hm_carbon_pct": ("HM carbon (%)", "{:.3f}"),
    "hm_silicon_pct": ("HM silicon (%)", "{:.3f}"),
    "hm_sulphur_pct": ("HM sulphur (%)", "{:.3f}"),
    "hm_other_pct": ("HM other (%)", "{:.3f}"),
    "confirm_pellet_inputs": ("Pellet inputs confirmed", "{}"),
    "static_dataset_use_link": ("Dataset via DATA_URL", "{}"),
    "pci_override_on": ("PCI override", "{}"),
    "pci_override_kg": ("PCI override (kg/THM)", "{:,.1f}"),
    "transition_move_pct": ("Transition step (%/rung)", "{:.1f}"),
}

SLAG_SETTING_LABELS = {
    "enabled": "Use full slag balance",
    "slag_correction_factor": "Slag correction factor",
    "pi_loss_pct": "PI loss (%)",
    "fe_to_pig_iron_fraction": "Fe to PI fraction",
    "mn_recovery_pct": "Mn/Ti recovery (%)",
    "sulphur_gas_loss_pct": "S gas loss (%)",
    "alkali_to_slag_fraction": "Alkali to slag",
    "fe_to_feo_factor": "Fe to FeO",
    "mn_to_mno_factor": "Mn to MnO",
}

TABLE_INPUTS = (
    ("applied_ore_editor_df", "Ores",
     (("ore_name", "Ore"), ("selected", "Sel."), ("min_share_pct", "Min %"),
      ("max_share_pct", "Max %"), ("stock_mt", "Stock MT"), ("price_rs_per_mt", "Rs/MT"),
      ("moisture_pct", "H2O %"), ("fe_t_pct", "Fe %"), ("sio2_pct", "SiO2 %"),
      ("al2o3_pct", "Al2O3 %"), ("cao_pct", "CaO %"), ("mgo_pct", "MgO %"))),
    ("applied_flux_editor_df", "Flux",
     (("display_name", "Flux"), ("enabled", "On"), ("wet_qty_mt", "Wet MT"),
      ("cao_pct", "CaO %"), ("mgo_pct", "MgO %"), ("sio2_pct", "SiO2 %"),
      ("al2o3_pct", "Al2O3 %"), ("loi_pct", "LOI %"), ("price_rs_per_mt", "Rs/MT"))),
    ("applied_fuel_ash_editor_df", "Fuel ash",
     (("display_name", "Fuel"), ("enabled", "On"), ("rate_kg_per_thm", "kg/THM"),
      ("ash_pct", "Ash %"), ("sio2_pct", "SiO2 %"), ("al2o3_pct", "Al2O3 %"),
      ("moisture_pct", "H2O %"), ("vm_pct", "VM %"), ("price_rs_per_mt", "Rs/MT"))),
    ("applied_dust_editor_df", "BF gas dust",
     (("display_name", "Dust"), ("enabled", "On"), ("quantity_kg_per_charge", "kg/charge"),
      ("wet_qty_mt", "Wet MT"), ("fe_pct", "Fe %"), ("sio2_pct", "SiO2 %"),
      ("al2o3_pct", "Al2O3 %"), ("cao_pct", "CaO %"))),
)


# --- small helpers ---------------------------------------------------------------


def _num(value: Any) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _fmt(value: Any, spec: str = "{:,.2f}") -> str:
    if value is None:
        return "—"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    number = _num(value)
    if number is not None and "{:" in spec:
        try:
            return spec.format(number)
        except (ValueError, TypeError):
            return str(value)
    return str(value)


def _signed(value: float | None, spec: str = "{:+,.0f}") -> str:
    return "—" if value is None else spec.format(value).replace("-", "−")


def _label(suffix: str) -> str:
    return suffix.replace("_", " ").strip().capitalize()


def _when(iso: Any) -> str:
    if not iso:
        return "—"
    try:
        return datetime.fromisoformat(str(iso)).strftime("%d %b %Y, %H:%M IST")
    except ValueError:
        return str(iso)


# --- docx building blocks ----------------------------------------------------------


class _Doc:
    """python-docx with the house style: numbered sections, styled tables, charts."""

    def __init__(self) -> None:
        from docx import Document
        from docx.enum.section import WD_ORIENT  # noqa: F401 - imported for completeness
        from docx.shared import Mm, Pt, RGBColor

        self.Pt, self.RGB, self.Mm = Pt, RGBColor, Mm
        self.doc = Document()
        self.doc.core_properties.title = "Blend Mix Optimiser Snapshot Report"
        self.doc.core_properties.author = "Evonith Steel BF2"
        self.doc.core_properties.subject = "Blend optimisation snapshot and replay record"
        section = self.doc.sections[0]
        section.page_height, section.page_width = Mm(297), Mm(210)
        for side in ("left_margin", "right_margin"):
            setattr(section, side, Mm(17))
        section.top_margin, section.bottom_margin = Mm(16), Mm(16)
        self.width_in = (210 - 34) / 25.4

        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10)
        normal.paragraph_format.space_after = Pt(4)
        for level, size in ((1, 14), (2, 11.5)):
            style = styles[f"Heading {level}"]
            style.font.name = "Calibri"
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor.from_string(TEAL)
            style.paragraph_format.space_before = Pt(14 if level == 1 else 8)
            style.paragraph_format.space_after = Pt(4)
            style.paragraph_format.keep_with_next = True
        self._section_no = 0

    # text
    def section(self, title: str) -> None:
        self._section_no += 1
        self.doc.add_heading(f"{self._section_no}. {title}", level=1)

    def sub(self, title: str) -> None:
        self.doc.add_heading(title, level=2)

    def para(self, text: str = "", *, size: float | None = None, colour: str | None = None,
             bold: bool = False, italic: bool = False, align: str | None = None):
        p = self.doc.add_paragraph()
        if text:
            run = p.add_run(text)
            run.bold, run.italic = bold, italic
            if size:
                run.font.size = self.Pt(size)
            if colour:
                run.font.color.rgb = self.RGB.from_string(colour)
        if align == "right":
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        return p

    def note(self, text: str) -> None:
        self.para(text, size=8.5, colour=GREY_TEXT, italic=True)

    def bullets(self, items: list[str]) -> None:
        for item in items:
            p = self.doc.add_paragraph(style="List Bullet")
            self._rich(p, item)

    def _rich(self, paragraph, text: str) -> None:
        """``**bold**`` spans inside a sentence."""
        for i, chunk in enumerate(text.split("**")):
            if chunk:
                paragraph.add_run(chunk).bold = bool(i % 2)

    # tables
    @staticmethod
    def shade(cell, fill: str) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        tc_pr = cell._tc.get_or_add_tcPr()
        for old in tc_pr.findall(qn("w:shd")):
            tc_pr.remove(old)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)

    @staticmethod
    def borders(table, colour: str = "D1D5DB", size: int = 4, inner: bool = True) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        tbl_pr = table._tbl.tblPr
        el = OxmlElement("w:tblBorders")
        edges = ("top", "left", "bottom", "right") + (("insideH", "insideV") if inner else ())
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            node = OxmlElement(f"w:{edge}")
            if edge in edges:
                node.set(qn("w:val"), "single")
                node.set(qn("w:sz"), str(size))
                node.set(qn("w:color"), colour)
            else:
                node.set(qn("w:val"), "nil")
            el.append(node)
        tbl_pr.append(el)

    def _table_geometry(
        self, table, widths: list[float], *, repeat_header: bool = False
    ) -> None:
        """Write one fixed DXA geometry to the table grid and every cell.

        Setting ``cell.width`` alone leaves Word's ``tblGrid`` at equal widths,
        which makes wide appendix tables render with a crushed first column.
        Keeping ``tblW``, ``tblGrid`` and ``tcW`` identical also makes the
        layout stable across Word and headless renderers.
        """

        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        if len(widths) != len(table.columns):
            raise ValueError("Table width count must match the number of columns")
        twips = [max(1, round(float(width) * 1440)) for width in widths]
        tbl_pr = table._tbl.tblPr

        tbl_w = tbl_pr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            tbl_pr.insert(0, tbl_w)
        tbl_w.set(qn("w:type"), "dxa")
        tbl_w.set(qn("w:w"), str(sum(twips)))

        layout = tbl_pr.find(qn("w:tblLayout"))
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            tbl_pr.append(layout)
        layout.set(qn("w:type"), "fixed")

        grid_cols = list(table._tbl.tblGrid.gridCol_lst)
        for col, width in zip(grid_cols, twips):
            col.set(qn("w:w"), str(width))

        for row_no, row in enumerate(table.rows):
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:cantSplit")) is None:
                tr_pr.append(OxmlElement("w:cantSplit"))
            if repeat_header and row_no == 0 and tr_pr.find(qn("w:tblHeader")) is None:
                header = OxmlElement("w:tblHeader")
                header.set(qn("w:val"), "true")
                tr_pr.append(header)

            for cell, width in zip(row.cells, twips):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tc_pr = cell._tc.get_or_add_tcPr()
                tc_w = tc_pr.get_or_add_tcW()
                tc_w.set(qn("w:type"), "dxa")
                tc_w.set(qn("w:w"), str(width))
                tc_mar = tc_pr.find(qn("w:tcMar"))
                if tc_mar is None:
                    tc_mar = OxmlElement("w:tcMar")
                    tc_pr.append(tc_mar)
                for edge, margin in (("top", 55), ("left", 70), ("bottom", 55), ("right", 70)):
                    node = tc_mar.find(qn(f"w:{edge}"))
                    if node is None:
                        node = OxmlElement(f"w:{edge}")
                        tc_mar.append(node)
                    node.set(qn("w:w"), str(margin))
                    node.set(qn("w:type"), "dxa")

    def _cell_text(self, cell, text: str, *, bold=False, size=9.0, colour=None, align=None):
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = self.Pt(0)
        run = p.add_run(str(text))
        run.bold = bold
        run.font.size = self.Pt(size)
        if colour:
            run.font.color.rgb = self.RGB.from_string(colour)
        if align == "right":
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def table(self, header: list[str], rows: list[list[Any]], *, widths: list[float] | None = None,
              numeric_from: int = 1, size: float = 9.0, total_last: bool = False,
              status_col: int | None = None):
        """A house-style table. Columns from ``numeric_from`` are right-aligned.

        ``status_col`` colours cells whose text starts with ✓ (green) or ✗ (red).
        """

        table = self.doc.add_table(rows=1, cols=len(header))
        table.autofit = False
        self.borders(table)
        for j, text in enumerate(header):
            cell = table.rows[0].cells[j]
            self._cell_text(cell, text, bold=True, size=size, colour="FFFFFF",
                            align="right" if j >= numeric_from else None)
            self.shade(cell, TEAL)
        for i, row in enumerate(rows):
            cells = table.add_row().cells
            last = total_last and i == len(rows) - 1
            for j, value in enumerate(row):
                text = "" if value is None else str(value)
                self._cell_text(cells[j], text, bold=last, size=size,
                                align="right" if j >= numeric_from else None)
                if i % 2 == 1:
                    self.shade(cells[j], ZEBRA)
                if last:
                    self.shade(cells[j], TEAL_LIGHT)
                if status_col is not None and j == status_col:
                    if text.startswith("✓"):
                        self.shade(cells[j], GREEN_FILL)
                        self._cell_text(cells[j], text, bold=True, size=size, colour=GREEN_TEXT)
                    elif text.startswith("✗"):
                        self.shade(cells[j], RED_FILL)
                        self._cell_text(cells[j], text, bold=True, size=size, colour=RED_TEXT)
        geometry = widths or [self.width_in / len(header)] * len(header)
        self._table_geometry(table, geometry, repeat_header=True)
        self.doc.add_paragraph().paragraph_format.space_after = self.Pt(2)
        return table

    def kv(self, pairs: list[tuple[str, str]], columns: int = 2) -> None:
        """Label/value pairs laid out ``columns`` pairs per row."""

        rows = []
        for i in range(0, len(pairs), columns):
            chunk = pairs[i:i + columns]
            row: list[Any] = []
            for label, value in chunk:
                row += [label, value]
            row += ["", ""] * (columns - len(chunk))
            rows.append(row)
        header = ["Item", "Value"] * columns
        pair_w = self.width_in / columns
        self.table(header, rows, widths=[pair_w * 0.64, pair_w * 0.36] * columns,
                   numeric_from=99, size=8.5)

    def image(self, png: bytes | None, width_in: float | None = None) -> None:
        if not png:
            return
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches

        self.doc.add_picture(io.BytesIO(png), width=Inches(width_in or self.width_in))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    def kpi_tiles(self, tiles: list[tuple[str, str, str]]) -> None:
        """One row of tiles: (value, label, sub-line)."""

        from docx.shared import Inches

        table = self.doc.add_table(rows=1, cols=len(tiles))
        table.autofit = False
        self.borders(table, colour="FFFFFF", size=12)
        w = self.width_in / len(tiles)
        for cell, (value, label, sub) in zip(table.rows[0].cells, tiles):
            cell.width = Inches(w)
            self.shade(cell, TEAL_LIGHT)
            self._cell_text(cell, value, bold=True, size=15, colour=TEAL, align="center")
            p = cell.add_paragraph()
            p.alignment = 1
            p.paragraph_format.space_after = self.Pt(0)
            run = p.add_run(label)
            run.font.size = self.Pt(8)
            run.font.color.rgb = self.RGB.from_string("374151")
            if sub:
                p2 = cell.add_paragraph()
                p2.alignment = 1
                p2.paragraph_format.space_after = self.Pt(2)
                r2 = p2.add_run(sub)
                r2.font.size = self.Pt(7.5)
                r2.font.color.rgb = self.RGB.from_string(GREY_TEXT)
        self._table_geometry(table, [w] * len(tiles))
        self.doc.add_paragraph().paragraph_format.space_after = self.Pt(2)

    def header_footer(self, left: str, snapshot_id: str) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        section = self.doc.sections[0]
        hp = section.header.paragraphs[0]
        hp.text = ""
        run = hp.add_run(left)
        run.font.size = self.Pt(8)
        run.font.color.rgb = self.RGB.from_string(GREY_TEXT)

        fp = section.footer.paragraphs[0]
        fp.text = ""
        r = fp.add_run(f"Snapshot {snapshot_id}   ·   Page ")
        r.font.size = self.Pt(8)
        r.font.color.rgb = self.RGB.from_string(GREY_TEXT)
        for code in ("PAGE", None, "NUMPAGES"):
            if code is None:
                t = fp.add_run(" of ")
                t.font.size = self.Pt(8)
                t.font.color.rgb = self.RGB.from_string(GREY_TEXT)
                continue
            fld_run = fp.add_run()
            fld_run.font.size = self.Pt(8)
            begin = OxmlElement("w:fldChar")
            begin.set(qn("w:fldCharType"), "begin")
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = code
            end = OxmlElement("w:fldChar")
            end.set(qn("w:fldCharType"), "end")
            fld_run._r.append(begin)
            fld_run._r.append(instr)
            fld_run._r.append(end)

    def page_break(self) -> None:
        from docx.enum.text import WD_BREAK

        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def bytes(self) -> bytes:
        buf = io.BytesIO()
        self.doc.save(buf)
        return buf.getvalue()


# --- charts ----------------------------------------------------------------------


def _plt():
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 8.5,
                         "axes.spines.top": False, "axes.spines.right": False})
    return plt


def _png(fig) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, bbox_inches="tight")
    _plt().close(fig)
    return buf.getvalue()


def _donut(shares: list[tuple[str, float]], centre: str) -> bytes | None:
    items = [(k, v) for k, v in shares if v > 0.05]
    if not items:
        return None
    plt = _plt()
    fig, ax = plt.subplots(figsize=(3.3, 3.3))
    ax.pie([v for _, v in items], startangle=90, counterclock=False,
           colors=CHART_COLOURS[:len(items)],
           wedgeprops={"width": 0.38, "edgecolor": "white", "linewidth": 1.5},
           autopct=lambda p: f"{p:.1f}%" if p >= 4 else "", pctdistance=0.81,
           textprops={"fontsize": 7.5, "color": "white", "fontweight": "bold"})
    ax.text(0, 0, centre, ha="center", va="center", fontsize=9.5, fontweight="bold",
            color="#0F766E")
    ax.axis("equal")
    return _png(fig)


def _cost_bars(options: list[tuple[str, float, float, float]]) -> bytes | None:
    options = [o for o in options if all(v is not None for v in o[1:])]
    if not options:
        return None
    plt = _plt()
    fig, ax = plt.subplots(figsize=(6.6, 0.55 + 0.5 * len(options)))
    labels = [o[0] for o in options][::-1]
    ore = [o[1] for o in options][::-1]
    fuel = [o[2] for o in options][::-1]
    flux = [o[3] for o in options][::-1]
    ax.barh(labels, ore, color="#0F766E", label="Ore")
    ax.barh(labels, fuel, left=ore, color="#F59E0B", label="Fuel")
    ax.barh(labels, flux, left=[a + b for a, b in zip(ore, fuel)], color="#6366F1", label="Flux")
    for y, (a, b, c) in enumerate(zip(ore, fuel, flux)):
        ax.text(a + b + c, y, f"  {a + b + c:,.0f}", va="center", fontsize=8.5, fontweight="bold")
    ax.set_xlabel("Rs / THM")
    ax.legend(ncol=3, loc="lower center", bbox_to_anchor=(0.5, 1.0), frameon=False, fontsize=8)
    ax.set_xlim(0, max(a + b + c for a, b, c in zip(ore, fuel, flux)) * 1.14)
    return _png(fig)


def _coke_waterfall(anchor: float, terms: list[tuple[str, float]], final: float) -> bytes | None:
    if anchor is None or final is None:
        return None
    plt = _plt()
    steps = [("Model anchor", anchor, "base")]
    steps += [(label, delta, "delta") for label, delta in terms if abs(delta) > 1e-9]
    steps.append(("Reported coke", final, "base"))
    fig, ax = plt.subplots(figsize=(6.6, 2.5))
    running = anchor
    lows = []
    for i, (label, value, kind) in enumerate(steps):
        if kind == "base":
            ax.bar(i, value, color="#0F766E", width=0.6)
            ax.text(i, value, f"{value:,.1f}", ha="center", va="bottom", fontsize=8, fontweight="bold")
            lows.append(value)
        else:
            bottom = running if value >= 0 else running + value
            ax.bar(i, abs(value), bottom=bottom, width=0.6,
                   color="#EF4444" if value > 0 else "#22C55E")
            ax.text(i, running + max(value, 0), f"{value:+.2f}", ha="center", va="bottom", fontsize=8)
            running += value
            lows.append(bottom)
    ax.set_xticks(range(len(steps)))
    ax.set_xticklabels([s[0] for s in steps], fontsize=7.5)
    ax.set_ylabel("kg / THM")
    lo = min(lows + [anchor, final])
    hi = max(anchor, final, running)
    pad = max(2.0, (hi - lo) * 0.6)
    ax.set_ylim(lo - pad, hi + pad * 0.8)
    return _png(fig)


def _slag_window(rows: list[tuple[str, float, float | None, float | None, str]]) -> bytes | None:
    """Bullet chart: each metric against its allowed band."""

    rows = [r for r in rows if r[1] is not None]
    if not rows:
        return None
    plt = _plt()
    fig, axes = plt.subplots(len(rows), 1, figsize=(6.6, 0.52 * len(rows) + 0.2))
    if len(rows) == 1:
        axes = [axes]
    for ax, (label, value, lo, hi, spec) in zip(axes, rows):
        points = [v for v in (value, lo, hi) if v is not None]
        span = (max(points) - min(points)) or abs(value) * 0.1 or 1.0
        x0, x1 = min(points) - span * 0.6, max(points) + span * 0.6
        ax.set_xlim(x0, x1)
        ax.set_ylim(0, 1)
        ax.axhspan(0.35, 0.65, color="#E5E7EB")
        band_lo = lo if lo is not None else x0
        band_hi = hi if hi is not None else x1
        ax.axvspan(band_lo, band_hi, ymin=0.35, ymax=0.65, color="#A7F3D0")
        ok = (lo is None or value >= lo - 1e-9) and (hi is None or value <= hi + 1e-9)
        ax.plot([value], [0.5], marker="D", markersize=7, color="#065F46" if ok else "#DC2626")
        ax.text(value, 0.85, spec.format(value), ha="center", va="center", fontsize=7.5,
                fontweight="bold", color="#065F46" if ok else "#DC2626")
        for bound in (lo, hi):
            if bound is not None:
                ax.text(bound, 0.12, spec.format(bound), ha="center", va="center",
                        fontsize=6.5, color="#6B7280")
        ax.set_yticks([])
        ax.set_xticks([])
        for spine in ax.spines.values():
            spine.set_visible(False)
        ax.text(x0, 0.5, f"{label}  ", ha="right", va="center", fontsize=8)
    fig.subplots_adjust(left=0.24, right=0.98, hspace=0.25)
    return _png(fig)


# --- reading the snapshot -----------------------------------------------------------


def _metrics(fields: Mapping[str, Any]) -> dict[str, Any]:
    """The numbers a report needs from one BlendEvaluation (decoded to a dict)."""

    diag = fields.get("diagnostics") or {}
    rates = diag.get("fuel_rate_estimate") or {}
    flux = _num(diag.get("flux_cost_per_thm_rs")) or 0.0
    total = _num(diag.get("adjusted_objective_rs_per_thm"))
    if total is None:
        total = _num(fields.get("objective_rs_per_thm"))
    fuel = _num(diag.get("adjusted_fuel_cost_per_thm_rs"))
    if fuel is None:
        fuel = _num(fields.get("fuel_cost_per_thm_rs"))
    return {
        "total": (total + flux) if total is not None else None,
        "ore": _num(fields.get("ore_cost_per_thm_rs")),
        "fuel": fuel,
        "flux": flux,
        "coke": _num(rates.get("coke_rate_kg_thm")),
        "nut": _num(rates.get("nut_coke_rate_kg_thm")),
        "pci": _num(rates.get("pci_rate_kg_thm")),
        "fuel_rate": _num(rates.get("total_fuel_rate_kg_thm")),
        "correction": _num(diag.get("coke_correction_delta_kg_thm")),
        "fe": _num(fields.get("fe_t_pct")),
        "slag_rate": _num(fields.get("slag_rate_kg_per_thm")),
        "slag_mt": _num(fields.get("slag_mt")),
        "b2": _num(fields.get("slag_basicity")),
        "tb": _num(fields.get("slag_t_basicity")),
        "ib4": _num(fields.get("slag_ib4")),
        "al2o3": _num(fields.get("slag_al2o3_pct")),
        "mgo": _num(fields.get("slag_mgo_pct")),
        "mgo_al2o3": _num(fields.get("slag_mgo_al2o3_ratio")),
        "burden_mt": _num(diag.get("total_burden_qty_mt")) or _num(fields.get("total_qty_mt")),
        "fe_mt": _num(fields.get("fe_production_mt")),
        "feasible": fields.get("feasible"),
        "violations": list(fields.get("violations") or []),
    }


def _as_fields(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _ore_rows(snapshot: Mapping[str, Any]) -> dict[str, dict[str, Any]]:
    rows = _frame_rows((snapshot.get("inputs") or {}).get("applied_ore_editor_df"))
    return {str(r.get("ore_id")): r for r in rows if r.get("ore_id") is not None}


def _history_end(snapshot: Mapping[str, Any]) -> tuple[str, int] | None:
    for key, value in ((snapshot.get("frozen") or {}).get("provider_calls") or {}).items():
        if not key.startswith("get_history_frame("):
            continue
        items = value.get("items") if isinstance(value, dict) else None
        frame = items[0] if items else None
        if isinstance(frame, dict) and frame.get("__type__") == "compact_frame":
            index = ((frame.get("rows") or {}).get("index") or {}).get("values") or []
            stamps = [v for v in index if v]
            return (stamps[-1] if stamps else "", int(frame.get("n_rows") or 0))
    return None


# --- the document ----------------------------------------------------------------


def build_docx(snapshot: Mapping[str, Any]) -> bytes:
    """Render a snapshot as a Word document and return the file's bytes."""

    d = _Doc()
    summary = snapshot.get("summary") or {}
    results = {k: decode(v, typed=False) for k, v in (snapshot.get("results") or {}).items()}
    context = {k: decode(v, typed=False) for k, v in (snapshot.get("context") or {}).items()}
    inputs = snapshot.get("inputs") or {}
    frozen = snapshot.get("frozen") or {}
    run = snapshot.get("run") or {}
    names = ore_names(snapshot)
    ore_rows = _ore_rows(snapshot)
    basis, _encoded_rec = recommended_result(snapshot)
    lp = _as_fields(results.get("lp_result"))
    de = _as_fields(results.get("de_result"))
    rec = de if basis.startswith("DE") else lp
    rec_m = _metrics(rec) if rec else {}
    rec_si = results.get("de_si") if basis.startswith("DE") else results.get("lp_si")
    manual_blend = _as_fields(results.get("manual_blend"))
    manual_q = results.get("manual_quantities_mt") or {}
    snapshot_id = str(snapshot.get("id") or "(unsaved)")

    # Compare against the current blend when the build evaluated it, else the LP.
    compare_label, compare_m = None, None
    if manual_blend:
        compare_label, compare_m = "current blend", _metrics(manual_blend)
    elif rec is de and lp and de:
        compare_label, compare_m = "LP baseline", _metrics(lp)

    d.header_footer("Evonith Steel · BF2 · Blend Mix Optimiser — Snapshot Report", snapshot_id)

    # --- cover -----------------------------------------------------------------------
    d.para("Blend Mix Optimiser", size=22, bold=True, colour=TEAL)
    d.para("Snapshot report" + (f" — {snapshot.get('label')}" if snapshot.get("label") else ""),
           size=13, colour="374151")
    frozen_ok = bool(frozen.get("provider_calls"))
    meta = [
        ("Run at", _when(run.get("at")) if run.get("at") else "—"),
        ("Snapshot taken", _when(snapshot.get("created_at"))),
        ("Taken on", str(snapshot.get("source") or "—")),
        ("Recommendation", basis or "No optimiser result recorded"),
        ("Plant data", "Frozen with the snapshot" if frozen_ok else "Not saved (earlier snapshot)"),
        ("Snapshot ID", snapshot_id),
    ]
    d.kv(meta, columns=2)

    if rec_m:
        def delta(key: str, spec: str, unit: str) -> str:
            if not compare_m or rec_m.get(key) is None or compare_m.get(key) is None:
                return ""
            return f"{_signed(rec_m[key] - compare_m[key], spec)} {unit} vs {compare_label}"

        d.kpi_tiles([
            (_fmt(rec_m["total"], "{:,.0f}"), "Total cost, Rs/THM", delta("total", "{:+,.0f}", "Rs")),
            (_fmt(rec_m["coke"], "{:,.1f}"), "Coke rate, kg/THM", delta("coke", "{:+,.1f}", "kg")),
            (_fmt(rec_m["fuel_rate"], "{:,.1f}"), "Fuel rate, kg/THM", delta("fuel_rate", "{:+,.1f}", "kg")),
            (_fmt(rec_m["slag_rate"], "{:,.0f}"), "Slag rate, kg/THM", delta("slag_rate", "{:+,.1f}", "kg")),
            (_fmt(rec_m["b2"], "{:.3f}"), "Basicity B2", delta("b2", "{:+.3f}", "")),
            (_fmt(rec_si, "{:.2f}"), "Predicted HM Si, %", ""),
        ])

    # --- executive summary --------------------------------------------------------------
    d.section("Executive summary")
    points: list[str] = []
    if rec:
        shares = sorted(((names.get(k, k), _num(v) or 0.0) for k, v in (rec.get("shares_pct") or {}).items()),
                        key=lambda kv: -kv[1])
        blend = ", ".join(f"{n} {v:.1f}%" for n, v in shares if v > 0.05)
        points.append(f"The **{basis}** optimiser recommends **{blend}**, at a total cost of "
                      f"**{_fmt(rec_m['total'], '{:,.0f}')} Rs/THM** for "
                      f"{_fmt(summary.get('production_mt'), '{:,.0f}')} MT hot metal.")
        if compare_m and compare_m.get("total") is not None and rec_m.get("total") is not None:
            diff = rec_m["total"] - compare_m["total"]
            word = "cheaper" if diff < 0 else "dearer"
            points.append(f"That is **{abs(diff):,.0f} Rs/THM {word}** than the {compare_label} "
                          f"({abs(diff) / compare_m['total'] * 100:.2f}%).")
        if rec_m.get("coke") is not None:
            corr = rec_m.get("correction")
            points.append(
                f"Coke rate **{rec_m['coke']:,.1f} kg/THM**, nut coke {_fmt(rec_m['nut'], '{:,.1f}')}, "
                f"PCI {_fmt(rec_m['pci'], '{:,.1f}')} — total fuel **{_fmt(rec_m['fuel_rate'], '{:,.1f}')} kg/THM**"
                + (f", including a physics correction of {_signed(corr, '{:+.2f}')} kg/THM." if corr else ".")
            )
        points.append(
            f"Slag **{_fmt(rec_m['slag_rate'], '{:,.1f}')} kg/THM** ({_fmt(rec_m['slag_mt'], '{:,.0f}')} MT), "
            f"B2 {_fmt(rec_m['b2'], '{:.3f}')}, T-basicity {_fmt(rec_m['tb'], '{:.3f}')}, "
            f"Al2O3 {_fmt(rec_m['al2o3'], '{:.2f}')}%, MgO {_fmt(rec_m['mgo'], '{:.2f}')}%."
        )
        if rec_m["violations"]:
            points.append(f"**{len(rec_m['violations'])} constraint(s) violated** — see the constraint check.")
        else:
            points.append("**Every constraint is met** — slag window, charging capacity and Fe target.")
        if de and (de.get("diagnostics") or {}).get("de_fell_back_to_lp"):
            points.append("The total-cost optimiser did not improve on the LP, so the LP blend is reported.")
    else:
        points.append("No optimiser result was on the page when this snapshot was taken; "
                      "it records the inputs only.")
    if not frozen_ok:
        points.append("This snapshot predates frozen plant data: replaying it uses the plant data "
                      "of the day it is replayed, so results can differ.")
    elif run.get("inputs_changed_after_run"):
        points.append("Inputs were edited after this run; the report shows the inputs the run used.")
    d.bullets(points)

    # --- recommended blend ----------------------------------------------------------------
    if rec:
        d.section(f"Recommended blend ({basis})")
        diag = rec.get("diagnostics") or {}
        quantities = rec.get("quantities_mt") or {}
        shares = rec.get("shares_pct") or {}
        dry = diag.get("dry_weight_mt_by_ore") or {}
        fe_mt = diag.get("fe_contribution_mt_by_ore") or {}
        slag_by = diag.get("slag_contribution_mt_by_ore") or {}
        order = sorted(shares, key=lambda k: -(_num(shares[k]) or 0))
        order = [k for k in order if (_num(shares[k]) or 0) > 1e-6]
        charge_mass = _num(inputs.get("charge_mass_mt"))
        charges = _num(diag.get("charge_count"))
        rows, totals = [], [0.0] * 5
        for ore_id in order:
            price = _num((ore_rows.get(ore_id) or {}).get("price_rs_per_mt"))
            wet = _num(quantities.get(ore_id)) or 0.0
            cost = wet * price / 1e5 if price is not None else None
            values = [wet, _num(dry.get(ore_id)), _num(fe_mt.get(ore_id)), _num(slag_by.get(ore_id)), cost]
            for i, v in enumerate(values):
                totals[i] += v or 0.0
            rows.append([
                names.get(ore_id, ore_id), _fmt(shares.get(ore_id), "{:.2f}"),
                _fmt(wet, "{:,.1f}"),
                _fmt(wet * 1000 / charges if charges else None, "{:,.0f}"),
                _fmt(values[1], "{:,.1f}"), _fmt(values[2], "{:,.1f}"), _fmt(values[3], "{:,.1f}"),
                _fmt(price, "{:,.0f}"), _fmt(cost, "{:,.2f}"),
            ])
        rows.append(["Total", "100.00", _fmt(totals[0], "{:,.1f}"),
                     _fmt(totals[0] * 1000 / charges if charges else None, "{:,.0f}"),
                     _fmt(totals[1], "{:,.1f}"), _fmt(totals[2], "{:,.1f}"),
                     _fmt(totals[3], "{:,.1f}"), "", _fmt(totals[4], "{:,.2f}")])
        png = _donut([(names.get(k, k), _num(shares[k]) or 0.0) for k in order],
                     f"{totals[0]:,.0f} MT\nburden")
        if png:
            from docx.shared import Inches

            holder = d.doc.add_table(rows=1, cols=2)
            d.borders(holder, colour="FFFFFF", inner=False)
            left, right = holder.rows[0].cells
            left.width, right.width = Inches(2.9), Inches(d.width_in - 2.9)
            d._table_geometry(holder, [2.9, d.width_in - 2.9])
            left.paragraphs[0].add_run().add_picture(io.BytesIO(png), width=Inches(2.7))
            for i, ore_id in enumerate(order):
                p = right.paragraphs[0] if i == 0 else right.add_paragraph()
                p.paragraph_format.space_after = d.Pt(2)
                swatch = p.add_run("■  ")
                swatch.font.color.rgb = d.RGB.from_string(CHART_COLOURS[i % len(CHART_COLOURS)][1:])
                swatch.font.size = d.Pt(11)
                name = p.add_run(f"{names.get(ore_id, ore_id)}  ")
                name.bold = True
                name.font.size = d.Pt(9.5)
                val = p.add_run(f"{_num(shares[ore_id]):.2f}%  ·  {_num(quantities.get(ore_id)) or 0:,.0f} MT")
                val.font.size = d.Pt(9)
            d.doc.add_paragraph()
        d.table(["Ore", "Share %", "Wet MT", "kg/charge", "Dry MT", "Fe MT", "Slag MT",
                 "Rs/MT", "₹ Lakhs"], rows,
                widths=[1.65, 0.62, 0.72, 0.72, 0.68, 0.62, 0.66, 0.62, 0.65],
                size=8.5, total_last=True)
        if charges or charge_mass:
            d.note(f"{_fmt(charges, '{:,.0f}')} charges/day at {_fmt(charge_mass, '{:.2f}')} MT each. "
                   "Slag MT is each ore's contribution before the hot-metal balance.")
        lp_flux = {k: v for k, v in (diag.get("lp_flux_quantities_mt") or {}).items() if (_num(v) or 0) > 1e-6}
        if lp_flux:
            d.sub("Flux added by the optimiser")
            d.table(["Flux", "Wet MT"], [[k, _fmt(v, "{:,.2f}")] for k, v in lp_flux.items()],
                    widths=[3.0, 1.2])

    # --- constraint check ---------------------------------------------------------------------
    if rec:
        d.section("Constraint check")
        tol = 1e-6
        checks: list[list[str]] = []

        def status(ok: bool | None) -> str:
            return "—" if ok is None else ("✓ Met" if ok else "✗ Violated")

        def row(name: str, limit: str, actual: str, ok: bool | None) -> None:
            checks.append([name, limit, actual, status(ok)])

        cap_mt = _num(context.get("target_slag_qty_mt"))
        if cap_mt is not None and rec_m.get("slag_mt") is not None:
            row("Slag quantity (model basis)", f"≤ {cap_mt:,.1f} MT", f"{rec_m['slag_mt']:,.1f} MT",
                rec_m["slag_mt"] <= cap_mt + tol)
        for label, key, lo_key, hi_key, spec in (
            ("Basicity B2 (CaO/SiO2)", "b2", "target_slag_basicity_min", "target_slag_basicity_max", "{:.3f}"),
            ("T-basicity (CaO+MgO)/SiO2", "tb", "target_slag_t_basicity_min", "target_slag_t_basicity_max", "{:.3f}"),
        ):
            lo, hi = _num(context.get(lo_key)), _num(context.get(hi_key))
            value = rec_m.get(key)
            if value is None or (lo is None and hi is None):
                continue
            limit = " – ".join(_fmt(v, spec) if v is not None else "open" for v in (lo, hi))
            row(label, limit, _fmt(value, spec),
                (lo is None or value >= lo - tol) and (hi is None or value <= hi + tol))
        for label, key, lim_key, sense, spec in (
            ("Slag Al2O3", "al2o3", "target_slag_al2o3_max_pct", "max", "{:.2f}%"),
            ("Slag MgO", "mgo", "target_slag_mgo_min_pct", "min", "{:.2f}%"),
            ("MgO / Al2O3", "mgo_al2o3", "target_slag_mgo_al2o3_ratio_min", "min", "{:.3f}"),
        ):
            limit = _num(context.get(lim_key))
            value = rec_m.get(key)
            if limit is None or value is None:
                continue
            ok = value <= limit + tol if sense == "max" else value >= limit - tol
            row(label, f"{'≤' if sense == 'max' else '≥'} {spec.format(limit)}", spec.format(value), ok)
        cap = _num(context.get("max_burden_qty_mt"))
        if cap is not None and rec_m.get("burden_mt") is not None:
            row("Burden within charging capacity", f"≤ {cap:,.0f} MT", f"{rec_m['burden_mt']:,.0f} MT",
                rec_m["burden_mt"] <= cap + tol)
        fe_target = _num(context.get("target_fe_mt"))
        if fe_target is not None and rec_m.get("fe_mt") is not None:
            row("Fe to hot metal", f"≥ {fe_target:,.1f} MT", f"{rec_m['fe_mt']:,.1f} MT",
                rec_m["fe_mt"] >= fe_target - 1e-3)
        verdict = "✓ Feasible" if rec_m.get("feasible") else f"✗ {len(rec_m['violations'])} violation(s)"
        checks.append(["Optimiser verdict", "", "", verdict])
        d.table(["Constraint", "Limit", "Result", "Status"], checks,
                widths=[2.5, 1.55, 1.55, 1.3], numeric_from=1, status_col=3)
        if rec_m["violations"]:
            d.sub("Violations reported by the optimiser")
            d.bullets([str(v) for v in rec_m["violations"]])

    # --- fuel and coke ------------------------------------------------------------------------
    if rec:
        d.section("Fuel and coke")
        diag = rec.get("diagnostics") or {}
        anchor = diag.get("fuel_rate_estimate_anchor") or {}
        correction = diag.get("coke_correction") or {}
        terms = [t for t in (correction.get("terms") or []) if isinstance(t, dict)]
        png = _coke_waterfall(
            _num(correction.get("anchor_coke_rate_kg_thm")) or _num(anchor.get("coke_rate_kg_thm")),
            [(str(t.get("label")), _num(t.get("delta_kg_thm")) or 0.0) for t in terms if t.get("enabled")],
            rec_m.get("coke"),
        )
        d.image(png, width_in=6.2)
        d.note("The model predicts the blend's fuel cost; coke is back-solved from it with nut coke "
               "and PCI held at their inputs (the anchor), then each physics term adjusts it.")
        prediction = diag.get("model_prediction") or {}
        rates_rows = [
            ["Coke", _fmt(anchor.get("coke_rate_kg_thm"), "{:,.2f}"), _fmt(rec_m.get("coke"), "{:,.2f}"),
             str(diag.get("fuel_rate_estimate_source") or "—")],
            ["Nut coke", _fmt(anchor.get("nut_coke_rate_kg_thm"), "{:,.2f}"), _fmt(rec_m.get("nut"), "{:,.2f}"),
             str(anchor.get("nut_coke_source") or "—")],
            ["PCI", _fmt(anchor.get("pci_rate_kg_thm"), "{:,.2f}"), _fmt(rec_m.get("pci"), "{:,.2f}"),
             str(anchor.get("pci_source") or "—")],
            ["Total fuel", _fmt(anchor.get("total_fuel_rate_kg_thm"), "{:,.2f}"),
             _fmt(rec_m.get("fuel_rate"), "{:,.2f}"), ""],
        ]
        d.table(["Fuel (kg/THM)", "Anchor", "Reported", "Source"], rates_rows,
                widths=[1.4, 1.0, 1.0, 3.5], numeric_from=1, total_last=True)
        if terms:
            d.sub("Physics correction terms")
            term_rows = []
            for t in terms:
                term_rows.append([
                    str(t.get("label")),
                    "On" if t.get("enabled") else "Off",
                    _fmt(t.get("x_blend"), "{:,.2f}"), _fmt(t.get("x_reference"), "{:,.2f}"),
                    str(t.get("x_units") or ""),
                    _signed(_num(t.get("delta_kg_thm")), "{:+.2f}") if t.get("enabled") else "—",
                    str(t.get("disabled_reason") or t.get("k_display") or ""),
                ])
            d.table(["Term", "State", "Blend", "Reference", "Units", "Δ coke", "Basis / note"],
                    term_rows, widths=[1.2, 0.5, 0.75, 0.8, 1.15, 0.6, 1.9], numeric_from=2, size=8)
        d.kv([
            ("Predicted fuel cost (model)", _fmt(prediction.get("value"), "{:,.2f} Rs/THM")),
            ("Model loaded", _fmt(prediction.get("model_loaded"))),
            ("Fallback formula used", _fmt(prediction.get("used_fallback"))),
            ("Predicted HM Si", _fmt(rec_si, "{:.3f} %")),
        ], columns=2)

    # --- slag -----------------------------------------------------------------------------------
    if rec:
        d.section("Slag")
        rows = [
            ("B2  CaO/SiO2", rec_m.get("b2"), _num(context.get("target_slag_basicity_min")),
             _num(context.get("target_slag_basicity_max")), "{:.3f}"),
            ("T-basicity", rec_m.get("tb"), _num(context.get("target_slag_t_basicity_min")),
             _num(context.get("target_slag_t_basicity_max")), "{:.3f}"),
            ("Al2O3 %", rec_m.get("al2o3"), None, _num(context.get("target_slag_al2o3_max_pct")), "{:.2f}"),
            ("MgO %", rec_m.get("mgo"), _num(context.get("target_slag_mgo_min_pct")), None, "{:.2f}"),
            ("MgO/Al2O3", rec_m.get("mgo_al2o3"), _num(context.get("target_slag_mgo_al2o3_ratio_min")), None, "{:.3f}"),
        ]
        d.image(_slag_window(rows), width_in=6.2)
        d.note("Green band: allowed window. Diamond: this blend (red when outside).")
        diag = rec.get("diagnostics") or {}
        denom = _num(diag.get("slag_chemistry_denominator_mt")) or rec_m.get("slag_mt")
        chem = []
        for label, key in (("SiO2", "slag_basicity_sio2_mt"), ("CaO", "slag_basicity_cao_mt"),
                           ("MgO", "slag_basicity_mgo_mt"), ("Al2O3", "slag_al2o3_mt")):
            mt = _num(diag.get(key))
            if mt is not None:
                chem.append([label, _fmt(mt, "{:,.1f}"), _fmt(mt / denom * 100 if denom else None, "{:.2f}")])
        sources = []
        for label, key in (("Ores", "ore_slag_mt"), ("Flux", "flux_slag_mt"), ("Fuel ash", "fuel_ash_slag_mt")):
            mt = _num(diag.get(key))
            if mt is not None:
                sources.append([label, _fmt(mt, "{:,.1f}")])
        sources.append(["Total slag", _fmt(rec_m.get("slag_mt"), "{:,.1f}")])
        if chem:
            d.sub("Slag chemistry")
            d.table(["Component", "MT", "% of slag"], chem, widths=[2.0, 1.3, 1.3])
        d.sub("Where the slag comes from")
        d.table(["Source", "MT"], sources, widths=[2.0, 1.3], total_last=True)
        d.note(f"Slag rate {_fmt(rec_m.get('slag_rate'), '{:,.1f}')} kg/THM on the model basis; "
               f"model-to-plant factor {_fmt(context.get('model_to_plant_slag_factor'), '{:.3f}')}. "
               f"Observed DPR slag rate {_fmt(context.get('observed_slag_rate'), '{:,.1f}')} kg/THM.")

    # --- cost and options -----------------------------------------------------------------------
    options: list[tuple[str, dict[str, Any], Any]] = []
    if manual_blend:
        options.append(("Current blend", _metrics(manual_blend), results.get("manual_si")))
    if lp:
        options.append(("LP baseline", _metrics(lp), results.get("lp_si")))
    if de:
        label = "DE total cost" + (" (= LP)" if (de.get("diagnostics") or {}).get("de_fell_back_to_lp") else "")
        options.append((label, _metrics(de), results.get("de_si")))
    if options:
        d.section("Cost and options compared")
        d.image(_cost_bars([(n, m.get("ore"), m.get("fuel"), m.get("flux")) for n, m, _ in options]),
                width_in=6.2)
        metric_rows = [
            ("Total cost (Rs/THM)", "total", "{:,.0f}"), ("Ore cost (Rs/THM)", "ore", "{:,.0f}"),
            ("Fuel cost (Rs/THM)", "fuel", "{:,.0f}"), ("Flux cost (Rs/THM)", "flux", "{:,.1f}"),
            ("Coke rate (kg/THM)", "coke", "{:,.2f}"), ("Total fuel (kg/THM)", "fuel_rate", "{:,.2f}"),
            ("Burden Fe (%)", "fe", "{:.2f}"), ("Slag rate (kg/THM)", "slag_rate", "{:,.1f}"),
            ("B2", "b2", "{:.3f}"), ("T-basicity", "tb", "{:.3f}"),
            ("Al2O3 (%)", "al2o3", "{:.2f}"), ("MgO (%)", "mgo", "{:.2f}"),
            ("Burden (MT)", "burden_mt", "{:,.0f}"),
        ]
        table_rows = [[label] + [_fmt(m.get(key), spec) for _, m, _ in options]
                      for label, key, spec in metric_rows]
        table_rows.append(["Predicted HM Si (%)"] + [_fmt(si, "{:.3f}") for _, _, si in options])
        table_rows.append(["Feasible"] + [_fmt(m.get("feasible")) for _, m, _ in options])
        col_w = (d.width_in - 2.2) / len(options)
        d.table(["Metric"] + [n for n, _, _ in options], table_rows,
                widths=[2.2] + [col_w] * len(options))
        if manual_q and not manual_blend:
            total_q = sum((_num(v) or 0.0) for v in manual_q.values()) or 1.0
            d.sub("Current blend (last shift)")
            d.table(["Ore", "Share %", "Wet MT"], [
                [names.get(k, k), _fmt((_num(v) or 0) / total_q * 100, "{:.2f}"), _fmt(v, "{:,.1f}")]
                for k, v in sorted(manual_q.items(), key=lambda kv: -(_num(kv[1]) or 0)) if (_num(v) or 0) > 0
            ], widths=[2.6, 1.0, 1.2])
            d.note("Quantities the last-shift blend needs for the same hot-metal target. "
                   "Its cost is not evaluated in this build.")

    # --- path and commentary (UAT builds) -----------------------------------------------------------
    process = results.get("process_recommendation") or decode(inputs.get("process_recommendation"), typed=False)
    ladder = results.get("transition_ladder") or decode(inputs.get("transition_ladder"), typed=False)
    if process or ladder:
        d.section("Process parameters and transition path")
        if isinstance(process, dict):
            d.kv([(_label(str(k)), _fmt(v)) for k, v in process.items()
                  if not isinstance(v, (dict, list))], columns=2)
        if isinstance(ladder, list) and ladder and isinstance(ladder[0], dict):
            cols = list(ladder[0].keys())[:8]
            d.table([_label(c) for c in cols], [[_fmt(r.get(c)) for c in cols] for r in ladder], size=8)
    commentary = results.get("commentary")
    if commentary:
        d.section("Furnace commentary")
        for block in str(commentary).split("\n\n"):
            text = block.strip()
            if text:
                d.para(text.replace("**", ""))

    # --- plant data ------------------------------------------------------------------------------
    d.section("Plant data used by this run")
    hm = context.get("hm_snapshot") if isinstance(context.get("hm_snapshot"), dict) else {}
    rates = context.get("recent_fuel_rates") if isinstance(context.get("recent_fuel_rates"), dict) else {}
    history = _history_end(snapshot)
    check = frozen.get("model_check") or {}
    source_rows = [
        ["Ore stock and chemistry", "Frozen" if frozen_ok else "Live at replay",
         f"mode {context.get('chemistry_mode', '—')}, window {context.get('chemistry_window_days', '—')} d"],
        ["Hot metal / slag", "Frozen" if frozen_ok else "Live at replay",
         f"HM Fe {_fmt(hm.get('hm_fe_pct_for_target'), '{:.2f}')}%, observed slag "
         f"{_fmt(hm.get('observed_slag_rate_kg_per_thm'), '{:,.1f}')} kg/THM"],
        ["Process history (coke / Si models)", "Frozen" if history else "Not saved",
         (f"{history[1]:,} hourly rows to {_when(history[0])}" if history else "—")],
        ["Live fuel rates (1 h)", "Frozen" if frozen_ok else "Live at replay",
         ", ".join(f"{k.replace('_rate_kg_thm', '')} {_fmt(v, '{:,.1f}')}"
                   for k, v in rates.items() if k.endswith("_kg_thm")) or "—"],
        ["Configuration (setting_bmo.yml)", "Frozen" if frozen_ok else "Live at replay", ""],
    ]
    d.table(["Source", "In snapshot", "Detail"], source_rows, widths=[2.2, 1.1, 3.6], numeric_from=99)
    if check:
        same = check.get("identical")
        d.para(("✓ " if same else "✗ ") + "Replay check: the fuel model on the saved history gives "
               f"{_fmt(check.get('model_output_saved_history'), '{:,.4f}')} Rs/THM against "
               f"{_fmt(check.get('model_output_full_history'), '{:,.4f}')} on the full history"
               + (" — identical." if same else " — NOT identical."),
               size=9, colour=GREEN_TEXT if same else RED_TEXT)

    # --- appendices ------------------------------------------------------------------------------
    d.page_break()
    d.section("Appendix — inputs")
    scalar_pairs = []
    for suffix, encoded in inputs.items():
        value = decode(encoded, typed=False)
        if isinstance(value, (dict, list)) or hasattr(value, "shape"):
            continue
        label, spec = INPUT_LABELS.get(suffix, (_label(suffix), "{}"))
        scalar_pairs.append((label, _fmt(value, spec)))
    if scalar_pairs:
        d.sub("Targets, limits and settings")
        d.kv(sorted(scalar_pairs), columns=2)
    slag_settings = context.get("slag_settings_values")
    if isinstance(slag_settings, dict) and slag_settings:
        d.sub("Advanced slag balance")
        d.kv([(SLAG_SETTING_LABELS.get(k, _label(k)), _fmt(v, "{:.4g}" if not isinstance(v, bool) else "{}"))
              for k, v in slag_settings.items()], columns=2)
    if context.get("de_seed_choice"):
        d.note(f"Total-cost optimiser start point: {context.get('de_seed_choice')}.")
    for suffix, title, columns in TABLE_INPUTS:
        rows = _frame_rows(inputs.get(suffix))
        if not rows:
            continue
        present = [(c, h) for c, h in columns if c in rows[0]]
        if suffix == "applied_ore_editor_df":
            rows = sorted(rows, key=lambda r: (not r.get("selected"), str(r.get("ore_name"))))
        body = []
        for r in rows:
            line = []
            for c, _h in present:
                v = decode(r.get(c), typed=False)
                line.append(_fmt(v, "{:,.2f}") if isinstance(v, float) else _fmt(v))
            body.append(line)
        d.sub(f"{title} ({len(rows)} rows)")
        first_w = 1.6
        rest = (d.width_in - first_w) / max(1, len(present) - 1)
        d.table([h for _c, h in present], body, widths=[first_w] + [rest] * (len(present) - 1),
                size=7.5)

    messages = []
    for key in ("lp_errors", "de_errors"):
        for message in results.get(key) or []:
            messages.append(f"{key.split('_')[0].upper()}: {message}")
    if messages:
        d.section("Appendix — optimiser messages")
        d.bullets(messages)

    d.section("Appendix — provenance and replay")
    prov = snapshot.get("provenance") or {}
    d.kv([
        ("Branch", str(prov.get("branch") or "—")),
        ("Commit", str(prov.get("commit") or "—")),
        ("Schema", str(snapshot.get("schema") or "—")),
        ("Inputs restorable", str(len(inputs))),
        ("Not restorable", ", ".join(snapshot.get("not_restorable") or []) or "—"),
        ("Report generated", datetime.now().strftime("%d %b %Y, %H:%M")),
    ], columns=2)
    d.note("To replay: on the Blend Mix Optimiser page choose this snapshot and click "
           "“Open in Sandbox” (or switch on Sandbox and load it). With frozen plant data, "
           "running LP/DE there reproduces the numbers in this report exactly.")
    return d.bytes()
