"""Document text extraction parsers for the Knowledge Hub upload pipeline.

Each function accepts a file-like object and returns the extracted text as a
plain Python string.  Supported formats:

* PDF — via PyMuPDF (``fitz``)
* DOCX — via python-docx
* PPTX — via python-pptx
* Excel (XLS/XLSX) — via pandas, rendered as Markdown table
"""

from io import BytesIO

import docx
import fitz  # PyMuPDF
import pandas as pd
import pptx


def parse_pdf(file) -> str:
    """Extract all text from a PDF file.

    Args:
        file: File-like object (binary mode) containing the PDF.

    Returns:
        Concatenated text from all pages, joined with newlines.
    """
    doc = fitz.open(stream=file.read(), filetype="pdf")
    return "\n".join([page.get_text() for page in doc])


def parse_docx(file) -> str:
    """Extract paragraph text from a DOCX file.

    Args:
        file: File-like object (binary mode) containing the DOCX.

    Returns:
        Paragraph texts joined with newlines.
    """
    doc = docx.Document(file)
    return "\n".join([para.text for para in doc.paragraphs])


def parse_pptx(file) -> str:
    """Extract shape text from all slides in a PPTX file.

    Args:
        file: File-like object (binary mode) containing the PPTX.

    Returns:
        All shape text joined with newlines.
    """
    prs = pptx.Presentation(file)
    text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text.append(shape.text)
    return "\n".join(text)


def parse_excel(file) -> str:
    """Read an Excel file and render the first sheet as a Markdown table.

    Args:
        file: File-like object (binary mode) containing the XLS/XLSX.

    Returns:
        Markdown-formatted table string.
    """
    df = pd.read_excel(file)
    return df.to_markdown()
